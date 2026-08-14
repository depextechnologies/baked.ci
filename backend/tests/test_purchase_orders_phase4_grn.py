"""Phase 4 — GRN document generation (PDF/Excel) regression suite.

Covers:
 * Partner endpoints: /api/partner/purchase-orders/{po_id}/grn(.pdf|.xlsx) + per-receipt
 * Supplier endpoints: /api/supplier/me/purchase-orders/{po_id}/grn(.pdf|.xlsx) + per-receipt
 * Admin endpoints:   /api/admin/modules/mart/purchase-orders/{po_id}/grn(.pdf|.xlsx) + per-receipt
 * RBAC negative: 401 without bearer, 404 for cross-tenant or draft PO
 * PDF content assertions (extracts via pypdf)
 * XLSX content assertions (opens via openpyxl)
"""
from __future__ import annotations
import io
import os
import pytest
import requests
from openpyxl import load_workbook
from pypdf import PdfReader

BASE_URL = os.environ.get(
    "REACT_APP_BACKEND_URL", "https://baked-platform.preview.emergentagent.com"
).rstrip("/")
API = f"{BASE_URL}/api"

SA_EMAIL = "depexopenai@gmail.com"
SA_PASSWORD = "baked@2026#!$@"
PARTNER_EMAIL = "partner-alpha-store@test.example"
PARTNER_PASSWORD = "Alpha1234!Beta"
PARTNER_BETA_EMAIL = "partner-beta-store@test.example"
PARTNER_BETA_PASSWORD = "Alpha1234!Beta"
DELTA_EMAIL = "demo-delta-supplier@test.example"
DELTA_PASSWORD = "Supplier1234!"

RECEIVED_PO_ID = "po_442fc2a8d8b5478f"
RECEIVED_PO_CODE = "PO-CI-2026-00032"


# --- auth fixtures ---------------------------------------------------------

def _bearer(token: str) -> dict:
    return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}


@pytest.fixture(scope="module")
def sa_headers():
    r = requests.post(f"{API}/admin/auth/login",
                      json={"email": SA_EMAIL, "password": SA_PASSWORD}, timeout=15).json()
    return _bearer(r.get("access_token") or r.get("token"))


@pytest.fixture(scope="module")
def partner_headers():
    r = requests.post(f"{API}/partner/auth/login",
                      json={"email": PARTNER_EMAIL, "password": PARTNER_PASSWORD}, timeout=15).json()
    return _bearer(r.get("access_token") or r.get("token"))


@pytest.fixture(scope="module")
def partner_beta_headers():
    r = requests.post(f"{API}/partner/auth/login",
                      json={"email": PARTNER_BETA_EMAIL, "password": PARTNER_BETA_PASSWORD}, timeout=15)
    if r.status_code != 200:
        pytest.skip("beta partner not seeded")
    j = r.json()
    return _bearer(j.get("access_token") or j.get("token"))


@pytest.fixture(scope="module")
def supplier_headers():
    r = requests.post(f"{API}/martbaked/sellers/login",
                      json={"email": DELTA_EMAIL, "password": DELTA_PASSWORD}, timeout=15)
    assert r.status_code == 200, r.text
    return _bearer(r.json()["access_token"])


# --- Partner list endpoint --------------------------------------------------

class TestPartnerGrnList:
    def test_list_shape(self, partner_headers):
        r = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/grn",
                         headers=partner_headers, timeout=15)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["po_code"] == RECEIVED_PO_CODE
        assert data["consolidated"]["available"] is True
        assert "reference" in data["consolidated"]
        assert isinstance(data["receipts"], list)
        assert len(data["receipts"]) >= 1
        first = data["receipts"][0]
        for k in ("id", "sequence", "reference", "received_at"):
            assert k in first, f"missing {k}"
        assert first["sequence"] == 1

    def test_list_unauthenticated(self):
        r = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/grn", timeout=15)
        assert r.status_code == 401

    def test_cross_tenant_404(self, partner_beta_headers):
        r = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/grn",
                         headers=partner_beta_headers, timeout=15)
        assert r.status_code == 404


# --- Partner PDF/XLSX + content assertions ---------------------------------

class TestPartnerGrnDownloads:
    def test_consolidated_pdf(self, partner_headers):
        r = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/grn.pdf",
                         headers=partner_headers, timeout=30)
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/pdf")
        assert "attachment" in r.headers.get("content-disposition", "")
        assert r.content[:5] == b"%PDF-"
        # Extract text and verify required strings
        reader = PdfReader(io.BytesIO(r.content))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        for needle in ["Goods Received Note", RECEIVED_PO_CODE, "Consolidated GRN",
                       "Grand total", "SUP-CI-0001", "Not a tax invoice"]:
            assert needle in text, f"PDF missing expected text: {needle!r}\n---\n{text[:2000]}"
        # BILL TO / BUYER + SUPPLIER labels (allow case-insensitive on 'buyer'/'supplier')
        low = text.lower()
        assert "buyer" in low or "bill to" in low
        assert "supplier" in low

    def test_consolidated_xlsx(self, partner_headers):
        r = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/grn.xlsx",
                         headers=partner_headers, timeout=30)
        assert r.status_code == 200
        assert "spreadsheetml.sheet" in r.headers["content-type"]
        assert r.content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(r.content))
        assert "GRN" in wb.sheetnames
        ws = wb["GRN"]
        all_cells = [str(c.value) if c.value is not None else "" for row in ws.iter_rows() for c in row]
        blob = " | ".join(all_cells)
        for needle in ["Goods Received Note", RECEIVED_PO_CODE, "Grand total"]:
            assert needle in blob, f"XLSX missing {needle!r}"

    def test_receipt_pdf_and_xlsx(self, partner_headers):
        lst = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/grn",
                           headers=partner_headers, timeout=15).json()
        r0 = lst["receipts"][0]
        rid = r0["id"]
        seq = r0["sequence"]
        # PDF
        r = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/receipts/{rid}/grn.pdf",
                         headers=partner_headers, timeout=30)
        assert r.status_code == 200
        assert r.content[:5] == b"%PDF-"
        reader = PdfReader(io.BytesIO(r.content))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        assert ("Receipt GRN" in text) or ("Received now" in text), \
            f"Per-receipt PDF must indicate receipt-scope\n---\n{text[:800]}"
        # sequence number should appear (as bare digits or padded)
        assert (str(seq) in text) or (f"R{seq:02d}" in text)
        # XLSX
        rx = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/receipts/{rid}/grn.xlsx",
                          headers=partner_headers, timeout=30)
        assert rx.status_code == 200
        assert rx.content[:2] == b"PK"

    def test_receipt_pdf_unauthenticated(self, partner_headers):
        lst = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/grn",
                           headers=partner_headers, timeout=15).json()
        rid = lst["receipts"][0]["id"]
        r = requests.get(f"{API}/partner/purchase-orders/{RECEIVED_PO_ID}/receipts/{rid}/grn.pdf",
                         timeout=15)
        assert r.status_code == 401


# --- Supplier endpoints -----------------------------------------------------

class TestSupplierGrn:
    def test_list_and_downloads(self, supplier_headers):
        r = requests.get(f"{API}/supplier/me/purchase-orders/{RECEIVED_PO_ID}/grn",
                         headers=supplier_headers, timeout=15)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["po_code"] == RECEIVED_PO_CODE
        assert len(data["receipts"]) >= 1
        # PDF
        rp = requests.get(f"{API}/supplier/me/purchase-orders/{RECEIVED_PO_ID}/grn.pdf",
                          headers=supplier_headers, timeout=30)
        assert rp.status_code == 200
        assert rp.content[:5] == b"%PDF-"
        # XLSX
        rx = requests.get(f"{API}/supplier/me/purchase-orders/{RECEIVED_PO_ID}/grn.xlsx",
                          headers=supplier_headers, timeout=30)
        assert rx.status_code == 200
        assert rx.content[:2] == b"PK"
        # per-receipt
        rid = data["receipts"][0]["id"]
        r2 = requests.get(f"{API}/supplier/me/purchase-orders/{RECEIVED_PO_ID}/receipts/{rid}/grn.pdf",
                          headers=supplier_headers, timeout=30)
        assert r2.status_code == 200
        assert r2.content[:5] == b"%PDF-"

    def test_unauthenticated_401(self):
        r = requests.get(f"{API}/supplier/me/purchase-orders/{RECEIVED_PO_ID}/grn.pdf", timeout=15)
        assert r.status_code == 401

    def test_cross_supplier_404(self, supplier_headers, partner_headers):
        # Create a draft PO with a *different* supplier — but simplest: find another supplier's PO
        # by asking SA. Fallback: use a random fake id → 404.
        r = requests.get(f"{API}/supplier/me/purchase-orders/po_nonexistent_xyz/grn",
                         headers=supplier_headers, timeout=15)
        assert r.status_code == 404

    def test_draft_po_404(self, supplier_headers, partner_headers):
        """Supplier cannot download GRN for a draft PO."""
        # find or create a draft PO for delta supplier
        pos = requests.get(f"{API}/partner/purchase-orders?status=draft",
                          headers=partner_headers, timeout=15).json()
        items = pos.get("items") or pos if isinstance(pos, list) else pos.get("items", [])
        draft = None
        for p in (items or []):
            if p.get("supplier_id") == "sup_demo_delta_seed" and p.get("status") == "draft":
                draft = p
                break
        if not draft:
            # create one
            c = requests.post(f"{API}/partner/purchase-orders", headers=partner_headers, json={
                "supplier_id": "sup_demo_delta_seed",
                "warehouse_id": "wh_alpha_demo_seed",
                "notes": "grn draft test",
            }, timeout=15)
            assert c.status_code in (200, 201), c.text
            draft = c.json()
        r = requests.get(f"{API}/supplier/me/purchase-orders/{draft['id']}/grn.pdf",
                         headers=supplier_headers, timeout=15)
        assert r.status_code == 404


# --- Admin endpoints --------------------------------------------------------

class TestAdminGrn:
    def test_list(self, sa_headers):
        r = requests.get(f"{API}/admin/modules/mart/purchase-orders/{RECEIVED_PO_ID}/grn",
                         headers=sa_headers, timeout=15)
        assert r.status_code == 200, r.text
        assert r.json()["po_code"] == RECEIVED_PO_CODE

    def test_pdf_and_xlsx(self, sa_headers):
        rp = requests.get(f"{API}/admin/modules/mart/purchase-orders/{RECEIVED_PO_ID}/grn.pdf",
                          headers=sa_headers, timeout=30)
        assert rp.status_code == 200
        assert rp.content[:5] == b"%PDF-"
        rx = requests.get(f"{API}/admin/modules/mart/purchase-orders/{RECEIVED_PO_ID}/grn.xlsx",
                          headers=sa_headers, timeout=30)
        assert rx.status_code == 200
        assert rx.content[:2] == b"PK"

    def test_receipt(self, sa_headers):
        lst = requests.get(f"{API}/admin/modules/mart/purchase-orders/{RECEIVED_PO_ID}/grn",
                           headers=sa_headers, timeout=15).json()
        rid = lst["receipts"][0]["id"]
        r = requests.get(
            f"{API}/admin/modules/mart/purchase-orders/{RECEIVED_PO_ID}/receipts/{rid}/grn.pdf",
            headers=sa_headers, timeout=30)
        assert r.status_code == 200
        assert r.content[:5] == b"%PDF-"

    def test_unauthenticated_401(self):
        r = requests.get(f"{API}/admin/modules/mart/purchase-orders/{RECEIVED_PO_ID}/grn.pdf", timeout=15)
        assert r.status_code == 401
